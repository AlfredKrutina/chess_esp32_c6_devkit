#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build CzechMate EO2026 PowerPoint (14 slides, 16:9)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "CzechMate_EO2026_prezentace.pptx"

# Geometry
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN

# Colors
CHARCOAL = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1E, 0x3A, 0x5F)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
PLACEHOLDER_BG = RGBColor(0xE8, 0xEA, 0xED)
PLACEHOLDER_BORDER = RGBColor(0xB0, 0xB4, 0xBA)
RULE = RGBColor(0xC8, 0xCD, 0xD4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_WHITE = RGBColor(0xFA, 0xFA, 0xFB)

TITLE_SIZE = Pt(26)
BODY_SIZE = Pt(16)
CAPTION_SIZE = Pt(12)
META_SIZE = Pt(14)


def set_run(run, size, color=CHARCOAL, bold=False, font_name="Calibri"):
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font_name


def add_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = NEAR_WHITE
    shape.line.fill.background()


def add_rule(slide, top=Inches(1.55)):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN, top, CONTENT_W, Emu(12700)  # ~1 pt
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.fill.background()


def add_title(slide, text, top=MARGIN):
    box = slide.shapes.add_textbox(MARGIN, top, CONTENT_W, Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    set_run(run, TITLE_SIZE, ACCENT, bold=True)
    return box


def add_body(slide, text, left=MARGIN, top=Inches(1.75), width=None, height=Inches(4.8)):
    if width is None:
        width = CONTENT_W
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    paragraphs = [p.strip() for p in text.strip().split("\n\n") if p.strip()]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = para
        set_run(run, BODY_SIZE, CHARCOAL, bold=False)
    return box


def add_placeholder(slide, left, top, width, height, caption):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PLACEHOLDER_BG
    shape.line.color.rgb = PLACEHOLDER_BORDER
    shape.line.width = Pt(1)

    cap = slide.shapes.add_textbox(left, top + height + Inches(0.08), width, Inches(0.35))
    tf = cap.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = caption
    set_run(run, CAPTION_SIZE, MUTED, bold=False)
    return shape


def add_footer(slide, page, total=14):
    box = slide.shapes.add_textbox(
        MARGIN, SLIDE_H - Inches(0.45), CONTENT_W, Inches(0.3)
    )
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{page} / {total}"
    set_run(run, Pt(11), MUTED)


def new_slide(prs):
    blank = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # --- 1 Title ---
    s = new_slide(prs)
    # accent bar left
    bar = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), SLIDE_H
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    box = s.shapes.add_textbox(MARGIN, Inches(1.8), CONTENT_W, Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "CzechMate: elektronická šachovnice s rozpoznáním figurek"
    set_run(run, Pt(32), ACCENT, bold=True)

    meta = s.shapes.add_textbox(MARGIN, Inches(3.4), CONTENT_W, Inches(2.2))
    tf = meta.text_frame
    tf.word_wrap = True
    lines = [
        ("Alfred Krutina", CHARCOAL, True, META_SIZE),
        ("Střední průmyslová škola elektrotechnická v Plzni", MUTED, False, META_SIZE),
        ("Elektrotechnická olympiáda EO2026", ACCENT, True, META_SIZE),
        ("", MUTED, False, META_SIZE),
        (
            "Šachovnice změří, která figurka na poli stojí, a ověří, zda tah smí.",
            CHARCOAL,
            False,
            BODY_SIZE,
        ),
    ]
    for i, (text, color, bold, size) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        if not text:
            continue
        run = p.add_run()
        run.text = text
        set_run(run, size, color, bold=bold)
    add_footer(s, 1)

    # --- 2 Deska ví, co na ní stojí ---
    s = new_slide(prs)
    add_title(s, "Hotová deska V2 ví, co na každém poli stojí.")
    add_rule(s)
    add_body(
        s,
        "CzechMate ve verzi V2 není jen podložka s diodami. Pod deskou měří Hallovy senzory "
        "magnetické pole figurek a firmware z toho určí typ i barvu. Na základě této informace "
        "kontroluje legálnost tahu a odpověď ukáže světlem na polích.\n\n"
        "Cílem je fyzická partie, u které deska sama drží pravidla — bez přepisování tahů do telefonu.",
        width=Inches(6.8),
        height=Inches(4.2),
    )
    add_placeholder(
        s,
        Inches(8.0),
        Inches(1.85),
        Inches(4.5),
        Inches(4.2),
        "[fotka: hotová deska V2 shora / šikmo]",
    )
    add_footer(s, 2)

    # --- 3 Kontaktní matice nestačí ---
    s = new_slide(prs)
    add_title(s, "Samotná kontaktní matice na rozpoznání figurek nestačí.")
    add_rule(s)
    add_body(
        s,
        "Reed spínač nebo kontaktní matice umí jen stav obsazeno / prázdné. Z toho nepoznáte, "
        "zda na poli stojí pěšec, nebo věž, ani jestli je figurka bílá, nebo černá. Partii proto "
        "nelze volně sestavit do libovolné pozice — deska by nevěděla, co hráč položil.\n\n"
        "Hallův senzor měří magnetické pole (sílu a polaritu). Když má každá figurka vlastní "
        "magnetickou stopu, deska získá typ i barvu, ne jen obsazenost.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 3)

    # --- 4 Dva Hallovy senzory ---
    s = new_slide(prs)
    add_title(s, "Každé pole má dva Hallovy senzory, aby šlo spolehlivě určit typ i barvu.")
    add_rule(s)
    add_body(
        s,
        "Dvojice senzorů dává diferenční měření vůči kalibrované nule. Z více vzorků se spočítá "
        "energie pole; prahová okna a hystereze (20 % u zamčeného typu) pocházejí z kalibrace "
        "na hotové desce. Debounce zabrání tomu, aby chvění při pokládání přepínalo typ.\n\n"
        "Výsledek pro každé pole je stabilní typ a barva figurky — ne surový výpis z osciloskopu.",
        width=Inches(6.8),
        height=Inches(4.2),
    )
    add_placeholder(
        s,
        Inches(8.0),
        Inches(1.85),
        Inches(4.5),
        Inches(4.2),
        "[schéma: jedno pole — 2× Hall, diff, klasifikace]",
    )
    add_footer(s, 4)

    # --- 5 Čtyři STM32 + ESP ---
    s = new_slide(prs)
    add_title(s, "Měření dělají čtyři STM32C031, pravidla a LED řídí ESP32-C6.")
    add_rule(s)
    add_body(
        s,
        "Osm krát osm polí je příliš mnoho kanálů na jeden mikrokontrolér s pohodlným ADC. "
        "Proto jsou čtyři STM32C031 — každý obsluhuje čtvrtinu šachovnice a přes I2C hlásí "
        "už rozpoznanou figurku (typ a barva), ne proud surových vzorků.\n\n"
        "ESP32-C6 drží šachová pravidla, stav partie a LED odpověď. Rozdělení je jasné: "
        "snímání na STM32, inteligence hry na ESP.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 5)

    # --- 6 Flash čtyř čipů ---
    s = new_slide(prs)
    add_title(s, "ESP32-C6 umí naprogramovat všechny čtyři STM32 ze stejné I2C sběrnice.")
    add_rule(s)
    add_body(
        s,
        "SDA, SCL a BOOT0 jsou sdílené; každý čip má vlastní NRST. Bootloader má adresu 0x63, "
        "proto smí být mimo reset vždy jen jeden STM32 — jinak by se na sběrnici srazily odpovědi. "
        "ESP drží ostatní v resetu, flashuje vybraný čip a pak přepne na další.\n\n"
        "Stejná sběrnice slouží později i pro běh: měření a aktualizace firmwaru bez výměny kabelů.",
        width=Inches(6.8),
        height=Inches(4.2),
    )
    add_placeholder(
        s,
        Inches(8.0),
        Inches(1.85),
        Inches(4.5),
        Inches(4.2),
        "[schéma: I2C — sdílené SDA/SCL/BOOT0, unikátní NRST]",
    )
    add_footer(s, 6)

    # --- 7 Jeden tah ---
    s = new_slide(prs)
    add_title(s, "Jeden tah na desce vede od zvednutí figurky až po přijetí nebo odmítnutí.")
    add_rule(s)
    add_body(
        s,
        "Hráč zvedne figurku — deska si pamatuje typ a barvu. Zobrazí zeleně legální cíle. "
        "Po položení firmware ověří, že na cíli stojí očekávaná figurka a že tah odpovídá "
        "pravidlům. Legální tah se přijme a stav partie se aktualizuje; nelegální se odmítne "
        "a světlo vyzve k opravě.\n\n"
        "Celý cyklus běží lokálně na ESP — bez nutnosti potvrzovat tah v aplikaci.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 7)

    # --- 8 Deska neschová chybu ---
    s = new_slide(prs)
    add_title(s, "Deska chybu neschová — nelegální tah, špatná figurka, šach i mat ukáže hned.")
    add_rule(s)
    add_body(
        s,
        "Když hráč položí figurku na pole, kam nesmí, deska tah nepřijme. Když na cíl dorazí "
        "jiný typ, než který se zvedl, systém to pozná z Hallů. Šach a mat se vyhodnocují "
        "ve stejném herním jádru jako legálnost tahu.\n\n"
        "Chyba tedy není tichá — hráč ji vidí na polích a musí ji fyzicky napravit, než partie pokračuje.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 8)

    # --- 9 Světlo ---
    s = new_slide(prs)
    add_title(s, "Světlo na polích je přímá odpověď desky na to, co hráč právě dělá.")
    add_rule(s)
    add_body(
        s,
        "Zelená typicky značí legální cíl, červená nebo výrazná chyba nelegální tah či špatnou "
        "figurku, další barvy vedou šach, mat a obnovu pozice. LED pásek bere samostatné "
        "externí 5 V, aby proud diod nezatěžoval logiku ESP a STM32.\n\n"
        "Hráč nemusí číst displej — stačí se dívat na desku.",
        width=Inches(6.8),
        height=Inches(4.2),
    )
    add_placeholder(
        s,
        Inches(8.0),
        Inches(1.85),
        Inches(4.5),
        Inches(4.2),
        "[fotka: LED na deskách při legálním tahu / chybě]",
    )
    add_footer(s, 9)

    # --- 10 Prohlížeč i telefon ---
    s = new_slide(prs)
    add_title(s, "Prohlížeč i telefon slouží jako pohled na partii, ne jako mozek desky.")
    add_rule(s)
    add_body(
        s,
        "Na ESP běží HTTP API a webové rozhraní — stav šachovnice a příkazy jdou přes síť. "
        "Flutter aplikace CzechMate ukazuje tutéž partii na telefonu: FEN, tahy, nastavení.\n\n"
        "Pravidla a rozpoznání figurek zůstávají ve firmwaru. Aplikace je displej a ovládání, "
        "ne náhrada elektroniky na stole.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 10)

    # --- 11 Tři problémy ---
    s = new_slide(prs)
    add_title(s, "Tři praktické problémy rozhodly, jak deska vypadá dnes.")
    add_rule(s)
    add_body(
        s,
        "Šum při snímání: prahová okna z laboratoře nestačila. Na hotové desce jsem doladil "
        "hysterezi a debounce — typ figurky se přestal přepínat při pokládání.\n\n"
        "Proud LED: desítky diod potřebují vlastní 5 V větev. Externí napájení drží jas bez "
        "poklesu napětí na logice.\n\n"
        "Flash čtyř STM32: společná bootloader adresa a omezená velikost firmwaru. Řešením je "
        "jeden čip mimo reset a úsporný kód na C031 — aktualizace běží z ESP po I2C.",
        top=Inches(1.75),
        height=Inches(5.0),
    )
    add_footer(s, 11)

    # --- 12 K čemu to je ---
    s = new_slide(prs)
    add_title(s, "CzechMate slouží ke hře a k učení šachů u fyzické šachovnice.")
    add_rule(s)
    add_body(
        s,
        "Začátečník vidí legální tahy přímo na polích a nemusí si pamatovat všechna pravidla "
        "najednou. Pokročilejší hráč dostane poctivou partii s kontrolou tahů bez přepisování "
        "do aplikace.\n\n"
        "Projekt je elektrotechnická a softwarová práce pro šachy a výuku — není to zdravotnická "
        "pomůcka ani nárok na asistivní zdravotní účel.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 12)

    # --- 13 Live demo ---
    s = new_slide(prs)
    add_title(s, "Ve finále zapnu tuhle desku a ukážu celý tah naživo.")
    add_rule(s)
    add_body(
        s,
        "Na stole bude hotová V2: zvednu figurku, ukážu zelené legální cíle, položím tah a deska "
        "ho přijme. Pak zkusím nelegální pole — světlo odmítne a vyzve k opravě. Krátce ukážu "
        "rozpoznání typu (např. pěšec versus věž) a stav v prohlížeči nebo telefonu.\n\n"
        "Cíl live části: rozhodčí uvidí snímání, pravidla i LED odpověď jako jeden celek.",
        top=Inches(1.85),
        height=Inches(4.5),
    )
    add_footer(s, 13)

    # --- 14 Závěr ---
    s = new_slide(prs)
    add_title(s, "CzechMate V2 je hotová elektronická šachovnice s rozpoznáním figurek.")
    add_rule(s)
    add_body(
        s,
        "Hallovy senzory a čtyři STM32 dodají typ a barvu; ESP32-C6 drží pravidla a světelnou odpověď.\n\n"
        "Deska umí přijmout legální tah a odmítnout chybu — bez přepisování partie do telefonu.\n\n"
        "Děkuji za pozornost. Rád odpovím na otázky k hardwaru, firmwaru i živé ukázce.",
        top=Inches(1.9),
        height=Inches(4.2),
    )
    add_footer(s, 14)

    prs.save(str(OUT))
    return OUT, len(prs.slides)


if __name__ == "__main__":
    path, n = build()
    print(f"Wrote {n} slides -> {path}")
