# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.child import INVALID_TITLE_REGEX
from copy import copy

black = Font(name="Calibri", size=11, color="FF000000")
bold = Font(name="Calibri", size=11, color="FF000000", bold=True)
linkf = Font(name="Calibri", size=11, color="FF0563C1", underline="single")
white = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")
thin = Border(
    left=Side(style="thin", color="FFBFBFBF"),
    right=Side(style="thin", color="FFBFBFBF"),
    top=Side(style="thin", color="FFBFBFBF"),
    bottom=Side(style="thin", color="FFBFBFBF"),
)
unlock = Protection(locked=False)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
mid = Alignment(horizontal="center", vertical="center", wrap_text=False)
num = Alignment(horizontal="right", vertical="center", wrap_text=False)

LINKS = {
    "10x1": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-10-mm-v-1-mm-sila-500-g_z42/",
    "10x2": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-10-mm-v-2-mm-sila-1-3-kg_z56/",
    "10x3": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-10-mm-v-3-mm-sila-1-8-kg_z113/",
    "10x5": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-10-mm-v-5-mm-sila-2-7-kg_z243/",
    "10x8": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-10-mm-v-8-mm-sila-3-7-kg_z331/",
    "12x6": "https://www.unimagnet.cz/neodymovy-magnet-kotouc-12-mm-v-6-mm-sila-4-5-kg_z45/",
}

# piece, disc, pull, bought, one set, energy, source, note, link key
PIECES = [
    ("pěšec", "Ø 10×1 mm, N42", "0,5 kg", 32, 16, 103050, "naměřeno",
     "Bez protiplechu je pod 150 000. Plech na zádech to může zvednout, číslo sem nepiš, dokud to nezměříš.", "10x1"),
    ("jezdec", "Ø 10×2 mm, N42", "1,3 kg", 8, 4, 324000, "odhad",
     "Odhad je asi 3× nad naměřeným pěšcem. Pěšec ale sám o sobě deska nemusí vidět.", "10x2"),
    ("střelec", "Ø 10×3 mm, N42", "1,8 kg", 8, 4, 583000, "odhad",
     "Odhad 1,8× nad jezdcem. Jeden kotouč místo tří tenkých.", "10x3"),
    ("věž", "Ø 10×5 mm, N38", "2,7 kg", 4, 4, 929000, "odhad",
     "Odhad jen 1,6× nad střelcem. Druhá sada chybí, jsou jen 4 kusy.", "10x5"),
    ("dáma", "Ø 10×8 mm, N42", "3,7 kg", 4, 2, 1685000, "odhad",
     "Odhad 1,8× nad věží. Tah 3,7 kg, figurky se budou cukat.", "10x8"),
    ("král", "Ø 12×6 mm, N42", "4,5 kg", 4, 2, 2098000, "odhad",
     "Odhad 1,25× nad dámou. S dámou se to může slít. Plech Ø 11 je menší než kotouč.", "12x6"),
]


def style(cell, font, align):
    cell.font = font
    cell.fill = white
    cell.alignment = align
    cell.border = thin
    cell.protection = unlock


def main():
    wb = openpyxl.load_workbook(r"tools/magnet_sila_cteni.xlsx")
    if "Koupit" in wb.sheetnames:
        del wb["Koupit"]
    k = wb.create_sheet("Koupit", 0)

    widths = [14, 22, 12, 12, 12, 12, 16, 14, 14, 12, 52, 62]
    for i, w in enumerate(widths, 1):
        k.column_dimensions[get_column_letter(i)].width = w

    k["A1"] = "Objednávka, jeden kotouč na figurku"
    k["A1"].font = bold
    k["A1"].fill = white
    k["A1"].alignment = left
    k.merge_cells("A1:L1")
    k.row_dimensions[1].height = 20

    k["A2"] = (
        "Koupeno jako dvě sady figurek (16 pěšců, 4 jezdci, 4 střelci, 4 věže, 2 dámy, 2 králové), "
        "věž jen jednou. Magnetů je 60, protiplechů 80, samolepek 70. "
        "Energie pěšce je měření bez plechu. Ostatní energie jsou odhad ze stejné mezery, taky bez plechu."
    )
    k["A2"].font = black
    k["A2"].fill = white
    k["A2"].alignment = left
    k.merge_cells("A2:L2")
    k.row_dimensions[2].height = 48

    headers = ["Figurka", "Kotouč", "Tah", "Koupeno", "Sada", "Druhá sada",
               "Energie", "S plechem", "Poměr", "Zdroj", "Poznámka", "Odkaz"]
    for col, text in enumerate(headers, 1):
        cell = k.cell(4, col, text)
        style(cell, bold, Alignment(horizontal="center", vertical="center", wrap_text=True))
    k.row_dimensions[4].height = 32

    for i, (name, disc, pull, bought, one, energy, source, note, key) in enumerate(PIECES):
        r = 5 + i
        vals = [
            (1, name, left, None),
            (2, disc, left, None),
            (3, pull, mid, None),
            (4, bought, mid, "0"),
            (5, one, mid, "0"),
            (6, "ano" if bought >= one * 2 else "ne", mid, None),
            (7, energy, num, "#,##0"),
            (8, None, num, "#,##0"),
            (10, source, mid, None),
            (11, note, left, None),
        ]
        for col, value, align, fmt in vals:
            cell = k.cell(r, col, value)
            style(cell, black, align)
            if fmt:
                cell.number_format = fmt
        if r == 5:
            cell = k.cell(r, 9, "—")
            style(cell, black, mid)
        else:
            cell = k.cell(r, 9, f'=IF(OR(G{r}="",G{r-1}=""),"",IF(H{r}="",G{r},H{r})/IF(H{r-1}="",G{r-1},H{r-1}))')
            style(cell, black, num)
            cell.number_format = "0.00"
        link = k.cell(r, 12, LINKS[key])
        style(link, linkf, left)
        link.hyperlink = LINKS[key]
        k.row_dimensions[r].height = 36

    k["A12"] = "Příslušenství, jeden kus na magnet a rezerva"
    k["A12"].font = bold
    k["A12"].fill = white
    k["A12"].alignment = left
    k.merge_cells("A12:L12")
    k.row_dimensions[12].height = 20

    extras = [
        (13, "protiplech", "OP-11-01-Zn", 80, "Ocelový plech na záda magnetu, ne mezi magnet a desku. 60 magnetů, 20 navíc."),
        (14, "samolepka", "3M 468MP, Ø 8 mm", 70, "Menší než kotouč, drží střed. 60 magnetů, 10 navíc."),
    ]
    for r, kind, what, qty, note in extras:
        for col, value, align in (
            (1, kind, left),
            (2, what, left),
            (4, qty, mid),
            (11, note, left),
        ):
            cell = k.cell(r, col, value)
            style(cell, black, align)
            if col == 4:
                cell.number_format = "0"
        for col in (3, 5, 6, 7, 8, 9, 10, 12):
            style(k.cell(r, col, None), black, left)
        k.row_dimensions[r].height = 32

    k.sheet_view.showGridLines = False
    k.sheet_view.zoomScale = 110
    k.sheet_view.view = "normal"
    k.freeze_panes = "A5"
    k.page_setup.orientation = "landscape"
    k.page_setup.paperSize = k.PAPERSIZE_A4
    k.page_setup.fitToWidth = 1
    k.page_setup.fitToHeight = 1
    k.sheet_properties.pageSetUpPr.fitToPage = True
    k.protection.sheet = False
    k.auto_filter.ref = None
    k.print_area = "A1:L14"
    k.sheet_properties.tabColor = None
    k.oddHeader.left.text = ""
    k.oddFooter.right.text = ""

    mer = wb["Měření"]
    mer["A33"] = "Řádky 1–9 jsou stohy malých kotoučů. Objednané jednotlivé kotouče jsou na listu Koupit."
    mer["A33"].font = black
    mer["A33"].fill = white

    wb.active = k
    saved = []
    for target in (r"tools/magnet_sila_cteni.xlsx", r"tools/magnet_sila.xlsx"):
        try:
            wb.save(target)
            saved.append(target)
        except PermissionError:
            saved.append(target + " LOCKED")
    print(" | ".join(saved))


if __name__ == "__main__":
    main()
